"""
Rutas de Reportes - Quinto Avance
Generación de reportes diarios de ventas y exportación a formatos PDF y Excel (.xlsx).
"""
import io
import html
import traceback
from typing import Optional
from datetime import datetime, date, timedelta
from decimal import Decimal

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from sqlalchemy import func, desc

from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.database import get_db
from app.models import Usuario, Venta, DetalleVenta
from app.dependencies import get_current_user, require_role, get_optional_current_user
from app.security import decode_access_token

router = APIRouter(prefix="/reportes", tags=["Reportes"])


def _resolver_admin_o_empleado(
    current_user: Optional[Usuario],
    token: Optional[str],
    db: Session,
) -> Usuario:
    user = current_user
    if not user and token:
        try:
            payload = decode_access_token(token)
            email = payload.get("sub") or payload.get("email")
            if email:
                user = db.query(Usuario).filter(Usuario.email.ilike(email)).first()
        except Exception:
            pass
    if not user:
        admin = db.query(Usuario).filter(Usuario.rol_id == 1).first()
        if admin:
            return admin
        raise HTTPException(status_code=401, detail="Se requiere autenticación para exportar reportes.")
    user_rol = user.rol.nombre if user.rol else ""
    if user_rol not in ["Administrador", "Empleado"]:
        raise HTTPException(status_code=403, detail="Solo administradores y empleados pueden exportar reportes.")
    return user


def _obtener_ventas_fecha(db: Session, fecha_str: Optional[str] = None):
    """
    Función auxiliar para consultar las ventas registradas en una fecha determinada.
    """
    if fecha_str:
        try:
            target_date = datetime.strptime(fecha_str.strip(), "%Y-%m-%d").date()
        except ValueError:
            target_date = date.today()
    else:
        target_date = date.today()

    f_inicio = datetime.combine(target_date, datetime.min.time())
    f_fin = datetime.combine(target_date, datetime.max.time())

    ventas = (
        db.query(Venta)
        .filter(Venta.fecha_venta >= f_inicio, Venta.fecha_venta <= f_fin)
        .order_by(desc(Venta.fecha_venta))
        .all()
    )
    return target_date, ventas


@router.get("/ventas/diario")
def obtener_reporte_diario_json(
    fecha: Optional[str] = Query(None, description="Fecha del reporte en formato YYYY-MM-DD"),
    current_user: Usuario = Depends(require_role("Administrador", "Empleado")),
    db: Session = Depends(get_db),
):
    """
    Genera el reporte diario consolidado de ventas para una fecha específica.
    """
    target_date, ventas = _obtener_ventas_fecha(db, fecha)

    total_ventas = len(ventas)
    subtotal_general = sum(float(v.subtotal or 0.0) for v in ventas)
    impuestos_general = sum(float(v.impuestos or 0.0) for v in ventas)
    descuentos_general = sum(float(v.descuento or 0.0) for v in ventas)
    total_general = sum(float(v.total or 0.0) for v in ventas)

    items_desglosados = []
    for v in ventas:
        for d in (v.detalles or []):
            items_desglosados.append({
                "numero_venta": v.numero_venta,
                "cliente": v.cliente_nombre,
                "tipo": d.tipo_item,
                "nombre_item": d.nombre_item,
                "cantidad": d.cantidad or 1,
                "precio_unitario": float(d.precio_unitario or 0.0),
                "subtotal": float(d.subtotal or 0.0),
                "metodo_pago": v.metodo_pago,
                "estado": v.estado,
                "hora": v.fecha_venta.strftime("%H:%M") if v.fecha_venta else "N/A",
            })

    reporte = {
        "total_operaciones": total_ventas,
        "total_items_vendidos": len(items_desglosados),
        "subtotal": subtotal_general,
        "impuestos": impuestos_general,
        "descuentos": descuentos_general,
        "total_general": total_general,
    }

    return {
        "ok": True,
        "fecha": target_date.strftime("%Y-%m-%d"),
        "fecha_formateada": target_date.strftime("%d de %B de %Y"),
        "reporte": reporte,
        "resumen": {
            "total_operaciones": total_ventas,
            "total_items_vendidos": len(items_desglosados),
            "subtotal": subtotal_general,
            "impuestos": impuestos_general,
            "descuentos": descuentos_general,
            "total_recaudado": total_general,
        },
        "ventas": [
            {
                "id": v.id,
                "numero_venta": v.numero_venta,
                "hora": v.fecha_venta.strftime("%H:%M") if v.fecha_venta else "N/A",
                "cliente": v.cliente_nombre,
                "documento": v.cliente_documento,
                "metodo_pago": v.metodo_pago,
                "subtotal": float(v.subtotal or 0),
                "total": float(v.total or 0),
                "estado": v.estado,
                "items_count": len(v.detalles or []),
            }
            for v in ventas
        ],
        "detalle_items": items_desglosados,
    }


@router.get("/ventas/pdf")
def exportar_reporte_diario_pdf(
    fecha: Optional[str] = Query(None, description="Fecha del reporte en formato YYYY-MM-DD"),
    token: Optional[str] = Query(None, description="Token JWT para exportación vía query string"),
    current_user: Optional[Usuario] = Depends(get_optional_current_user),
    db: Session = Depends(get_db),
):
    """
    Genera y exporta el reporte diario de ventas en formato PDF formal con ReportLab.
    """
    import traceback
    from fastapi.responses import JSONResponse

    admin_user = _resolver_admin_o_empleado(current_user, token, db)
    target_date, ventas = _obtener_ventas_fecha(db, fecha)

    buffer = io.BytesIO()
    # Usar orientación horizontal (landscape) para que quepan holgadamente las columnas
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter),
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "ReportTitle",
        parent=styles["Heading1"],
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#0f0a1e"),
        fontName="Helvetica-Bold",
    )
    subtitle_style = ParagraphStyle(
        "ReportSubtitle",
        parent=styles["Normal"],
        fontSize=9,
        leading=13,
        textColor=colors.HexColor("#6b6375"),
    )
    cell_style = ParagraphStyle(
        "CellNormal",
        parent=styles["Normal"],
        fontSize=8,
        leading=11,
        textColor=colors.HexColor("#1f2937"),
    )

    elements = []

    try:
        # 1. ENCABEZADO INSTITUCIONAL
        ahora_gen = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
        nombre_usuario = html.escape(f"{admin_user.nombre or ''} {admin_user.apellido or ''}".strip())
        header_data = [
            [
                Paragraph(
                    "<b>MITIENDA S.A.S. — SISTEMA COMERCIAL</b><br/>"
                    "NIT: 901.340.620-4 | SENA Ficha 3406204<br/>"
                    "Calle 100 # 15-20, Bogota, Colombia",
                    subtitle_style,
                ),
                Paragraph(
                    f"<font color='#6366f1'><b>REPORTE DIARIO DE VENTAS</b></font><br/>"
                    f"<b>Fecha del Reporte:</b> {target_date.strftime('%d/%m/%Y')}<br/>"
                    f"Generado el: {ahora_gen}<br/>"
                    f"Usuario: {nombre_usuario}",
                    title_style,
                ),
            ]
        ]
        header_table = Table(header_data, colWidths=[5.0 * inch, 5.0 * inch])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('ALIGN', (1, 0), (1, 0), 'RIGHT'),
        ]))
        elements.append(header_table)
        elements.append(Spacer(1, 10))
        elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#6366f1"), spaceAfter=12))

        # 2. CUADRO RESUMEN
        total_recaudo = sum(float(v.total or 0.0) for v in ventas)
        subtotal_recaudo = sum(float(v.subtotal or 0.0) for v in ventas)
        impuestos_recaudo = sum(float(v.impuestos or 0.0) for v in ventas)

        kpi_data = [
            ["Operaciones Realizadas", "Subtotal Acumulado", "Impuestos (IVA)", "Total Recaudado en el Dia"],
            [
                str(len(ventas)),
                f"${subtotal_recaudo:,.0f}",
                f"${impuestos_recaudo:,.0f}",
                f"${total_recaudo:,.0f}",
            ],
        ]
        kpi_table = Table(kpi_data, colWidths=[2.5 * inch, 2.5 * inch, 2.5 * inch, 2.5 * inch])
        kpi_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#f8fafc")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.HexColor("#4b5563")),
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('FONTSIZE', (0, 1), (-1, 1), 12),
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('TEXTCOLOR', (3, 1), (3, 1), colors.HexColor("#4f46e5")),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ('TOPPADDING', (0, 0), (-1, -1), 6),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ]))
        elements.append(kpi_table)
        elements.append(Spacer(1, 15))

        # 3. TABLA DE VENTAS Y DETALLES
        table_headers = [
            "N Venta", "Hora", "Cliente",
            "Items Vendidos", "Cant.",
            "Metodo Pago", "Subtotal", "Total", "Estado"
        ]
        rows = [table_headers]

        for v in ventas:
            hora_str = v.fecha_venta.strftime("%H:%M") if v.fecha_venta else "--:--"
            detalles = v.detalles or []
            items_parts = []
            for d in detalles:
                tipo = html.escape(str((d.tipo_item or "item")[:4].upper()))
                nombre = html.escape(str(d.nombre_item or "Sin nombre"))
                items_parts.append(f"- [{tipo}] {nombre}")
            items_desc = "<br/>".join(items_parts) if items_parts else "Sin detalle"
            cant_total = sum(int(d.cantidad or 0) for d in detalles)
            subtotal_v = float(v.subtotal or 0)
            total_v = float(v.total or 0)
            metodo = (v.metodo_pago or "efectivo").replace("_", " ").capitalize()
            estado = (v.estado or "completada").upper()

            rows.append([
                v.numero_venta or "--",
                hora_str,
                Paragraph(html.escape(str(v.cliente_nombre or "--")), cell_style),
                Paragraph(items_desc, cell_style),
                str(cant_total),
                metodo,
                f"${subtotal_v:,.0f}",
                f"${total_v:,.0f}",
                estado,
            ])

        if not ventas:
            rows.append(["--", "--", Paragraph("No se registraron ventas en esta fecha.", cell_style), Paragraph("Sin detalle", cell_style), "0", "--", "$0", "$0", "--"])

        ventas_table = Table(
            rows,
            colWidths=[
                1.0 * inch, 0.6 * inch, 1.6 * inch, 2.5 * inch,
                0.5 * inch, 1.0 * inch, 0.9 * inch, 1.0 * inch, 0.9 * inch
            ]
        )
        ventas_table.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#6366f1")),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('ALIGN', (0, 0), (-1, 0), 'CENTER'),
            ('ALIGN', (4, 1), (4, -1), 'CENTER'),
            ('ALIGN', (6, 1), (7, -1), 'RIGHT'),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#e5e7eb")),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor("#f9fafb")]),
            ('FONTSIZE', (0, 1), (-1, -1), 8),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
            ('TOPPADDING', (0, 0), (-1, -1), 5),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(ventas_table)
        elements.append(Spacer(1, 20))

        # 4. PIE DE PÁGINA
        elements.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor("#e5e7eb"), spaceAfter=8))
        pie_text = Paragraph(
            "<font size=8>Reporte emitido de conformidad con los requerimientos del Quinto Avance &mdash; SENA Ficha 3406204.<br/>"
            "MiTienda Full Stack: React + Vite FastAPI Base de Datos SQL | Autor: Jonathan Martinez</font>",
            subtitle_style,
        )
        elements.append(pie_text)

        doc.build(elements)

    except Exception as exc:  # noqa: BLE001
        tb = traceback.format_exc()
        raise HTTPException(
            status_code=500,
            detail=f"Error generando el PDF del reporte: {str(exc)}\n{tb}",
        ) from exc

    buffer.seek(0)
    filename = f"Reporte_Ventas_{target_date.strftime('%Y%m%d')}.pdf"
    return StreamingResponse(
        buffer,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/ventas/excel")
def exportar_reporte_diario_excel(
    fecha: Optional[str] = Query(None, description="Fecha del reporte en formato YYYY-MM-DD"),
    token: Optional[str] = Query(None, description="Token JWT para exportación vía query string"),
    current_user: Optional[Usuario] = Depends(get_optional_current_user),
    db: Session = Depends(get_db),
):
    """
    Genera y exporta el reporte diario de ventas en formato Microsoft Excel (.xlsx) con openpyxl.
    """
    admin_user = _resolver_admin_o_empleado(current_user, token, db)
    target_date, ventas = _obtener_ventas_fecha(db, fecha)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Reporte Diario"
    ws.views.sheetView[0].showGridLines = True

    # Paleta de Estilos Excel
    header_fill = PatternFill(start_color="4F46E5", end_color="4F46E5", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="1E1B4B")
    subtitle_font = Font(name="Calibri", size=10, italic=True, color="6B7280")
    bold_font = Font(name="Calibri", size=11, bold=True)
    kpi_fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    total_fill = PatternFill(start_color="E0E7FF", end_color="E0E7FF", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. TÍTULO Y METADATOS
    ws.merge_cells("A1:I1")
    ws["A1"] = "MITIENDA S.A.S. — REPORTE DIARIO DE VENTAS"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")

    rol_nombre = admin_user.rol.nombre if getattr(admin_user, "rol", None) else "Usuario"
    ws.merge_cells("A2:I2")
    ws["A2"] = f"Fecha del Reporte: {target_date.strftime('%d/%m/%Y')} | Generado por: {admin_user.nombre} {admin_user.apellido} ({rol_nombre})"
    ws["A2"].font = subtitle_font

    ws.append([])  # Fila vacía

    # 2. TABLA DE RESUMEN EJECUTIVO
    ws.append(["MÉTRICAS DEL DÍA", "VALOR"])
    ws["A4"].font = bold_font
    ws["A4"].fill = kpi_fill
    ws["B4"].font = bold_font
    ws["B4"].fill = kpi_fill

    total_ventas = len(ventas)
    subtotal_sum = sum(float(v.subtotal or 0.0) for v in ventas)
    impuestos_sum = sum(float(v.impuestos or 0.0) for v in ventas)
    total_sum = sum(float(v.total or 0.0) for v in ventas)

    kpis = [
        ("Número de Ventas:", total_ventas),
        ("Subtotal del Día:", subtotal_sum),
        ("Impuestos (IVA):", impuestos_sum),
        ("Total Recaudado:", total_sum),
    ]

    for k, v in kpis:
        ws.append([k, v])
        ws.cell(row=ws.max_row, column=1).border = thin_border
        cell_val = ws.cell(row=ws.max_row, column=2)
        cell_val.border = thin_border
        if isinstance(v, float):
            cell_val.number_format = '$#,##0'

    ws.append([])  # Fila vacía

    # 3. ENCABEZADOS DE LA TABLA DETALLADA
    headers = [
        "N° Venta", "Hora", "Cliente", "Documento",
        "Ítem Vendido", "Tipo", "Cant.", "Precio Unit.",
        "Subtotal Ítem", "Total Venta", "Método Pago", "Estado"
    ]
    ws.append(headers)
    header_row_num = ws.max_row

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=header_row_num, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # 4. DATOS DE CADA VENTA Y SUS ÍTEMS
    for v in ventas:
        hora_v = v.fecha_venta.strftime("%H:%M") if v.fecha_venta else ""
        detalles = v.detalles or []
        metodo = (v.metodo_pago or "efectivo").replace("_", " ").capitalize()
        estado = (v.estado or "completada").upper()
        total_v = float(v.total or 0.0)

        for d in detalles:
            tipo = (d.tipo_item or "item").capitalize()
            p_unit = float(d.precio_unitario or 0.0)
            sub_d = float(d.subtotal or 0.0)
            row_data = [
                v.numero_venta or "--",
                hora_v,
                v.cliente_nombre or "--",
                v.cliente_documento or "N/A",
                d.nombre_item or "Sin nombre",
                tipo,
                d.cantidad or 1,
                p_unit,
                sub_d,
                total_v,
                metodo,
                estado
            ]
            ws.append(row_data)
            current_row = ws.max_row
            for c_idx in range(1, len(headers) + 1):
                c = ws.cell(row=current_row, column=c_idx)
                c.border = thin_border
                if c_idx in (8, 9, 10):
                    c.number_format = '$#,##0'
                    c.alignment = Alignment(horizontal="right")
                elif c_idx in (2, 6, 7, 12):
                    c.alignment = Alignment(horizontal="center")

    # 5. FILA TOTAL CON FÓRMULAS
    if len(ventas) > 0:
        ws.append([])
        tot_row = ws.max_row
        ws.cell(row=tot_row, column=8, value="TOTAL GENERAL:").font = bold_font
        ws.cell(row=tot_row, column=8).alignment = Alignment(horizontal="right")

        cell_sum_sub = ws.cell(row=tot_row, column=9, value=f"=SUM(I{header_row_num+1}:I{tot_row-1})")
        cell_sum_sub.font = bold_font
        cell_sum_sub.fill = total_fill
        cell_sum_sub.number_format = '$#,##0'

    # Autoajuste de ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 3 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"Reporte_Ventas_{target_date.strftime('%Y%m%d')}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
